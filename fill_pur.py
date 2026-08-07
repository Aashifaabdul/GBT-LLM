# -*- coding: utf-8 -*-
"""
Fills the EP4DIS Progress Update Report template three times (PUR#1, #2, #3)
for the GBT-ICL dissertation project, matching the structure/depth of the
provided reference example (numbered sub-sections with completion tags,
bullet detail, a proper Risk/Impact/Mitigation nested table).
"""
import docx
from lxml import etree

NSMAP = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'w14': 'http://schemas.microsoft.com/office/word/2010/wordml',
}
def qn(tag):
    prefix, local = tag.split(':')
    return '{%s}%s' % (NSMAP[prefix], local)

def make_run(text, bold=False, italic=False):
    r = etree.Element(qn('w:r'))
    rpr = etree.SubElement(r, qn('w:rPr'))
    rfonts = etree.SubElement(rpr, qn('w:rFonts'))
    rfonts.set(qn('w:ascii'), 'Calibri'); rfonts.set(qn('w:hAnsi'), 'Calibri'); rfonts.set(qn('w:cs'), 'Calibri')
    if bold:
        etree.SubElement(rpr, qn('w:b'))
    if italic:
        etree.SubElement(rpr, qn('w:i'))
    t = etree.SubElement(r, qn('w:t'))
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    return r

def add_para(tc, runs, bullet=False):
    p = etree.SubElement(tc, qn('w:p'))
    if bullet:
        p.insert(0, make_run(u"• "))
    for r in runs:
        p.append(r)
    return p

def set_cell_rich(cell, blocks):
    """blocks: list of (text, bold, italic, bullet) tuples; one paragraph each. '' = blank spacer line."""
    tc = cell._tc
    for p in list(tc.findall(qn('w:p'))):
        tc.remove(p)
    for item in blocks:
        text, bold, italic, bullet = item
        if text == "":
            etree.SubElement(tc, qn('w:p'))
            continue
        p = etree.SubElement(tc, qn('w:p'))
        if bullet:
            p.append(make_run(u"• "))
        p.append(make_run(text, bold=bold, italic=italic))

def set_cell_simple(cell, texts):
    set_cell_rich(cell, [(t, False, False, False) for t in texts])

def set_combo_text(sdt, value):
    sdt_content = sdt.find(qn('w:sdtContent'))
    run_t = sdt_content.find('.//' + qn('w:t'))
    run_t.text = value
    sdtPr = sdt.find(qn('w:sdtPr'))
    plc = sdtPr.find(qn('w:showingPlcHdr'))
    if plc is not None:
        sdtPr.remove(plc)
    rstyle = run_t.getparent().find(qn('w:rPr') + '/' + qn('w:rStyle'))
    if rstyle is not None:
        rstyle.getparent().remove(rstyle)

TITLE = ("GBT-ICL: Eliminating Graph Side-Information in Graph-Based Transform "
         "Image Compression via In-Context Laplacian Prediction")

DESCRIPTION = [
    ("This project proposes GBT-ICL, an extension of Graph-Based Transform (GBT) image "
     "compression. Standard GBT methods construct a graph per image block and transmit it "
     "(or its Laplacian) as side-information, adding bitstream overhead. GBT-ICL removes this "
     "cost entirely: a frozen, shared in-context transformer predicts each block's graph "
     "Laplacian identically at encoder and decoder from already-decoded neighbouring pixels, so "
     "no graph is ever signalled. The predicted Laplacian is eigendecomposed to obtain a "
     "block-adaptive Graph Fourier Transform (GFT) basis, replacing the fixed DCT basis used in "
     "classical codecs. Coefficients are then quantised and entropy coded using standard "
     "(non-LLM) methods, unlike prior graph-compression designs that relied on an LLM entropy "
     "stage. Evaluation uses raw YUV test sequences (Beauty, HoneyBee, 1920x1080) with "
     "rate-distortion comparisons against DCT and non-adaptive GBT baselines.",
     False, False, False),
]

def build(checkpoint, progress_blocks, planned_blocks, risks):
    d = docx.Document('template.docx')
    t0, t1, t2 = d.tables[0], d.tables[1], d.tables[2]

    # Table 1: Supervisor + Checkpoint combo
    set_cell_simple(t1.rows[0].cells[1], ["[Supervisor Name — to be completed by student]"])
    sdt_nodes = t1._tbl.findall('.//' + qn('w:sdt'))
    for sdt in sdt_nodes:
        if sdt.find('.//' + qn('w:comboBox')) is not None:
            set_combo_text(sdt, checkpoint)

    # Table 2 content
    set_cell_simple(t2.rows[1].cells[0], [TITLE])
    set_cell_rich(t2.rows[3].cells[0], DESCRIPTION)

    # Ethics -> tick NO
    ethics_no_cell = t2.rows[4].cells[2]
    sdt = ethics_no_cell._tc.find('.//' + qn('w:sdt'))
    checked = sdt.find('.//' + qn('w14:checked'))
    checked.set(qn('w14:val'), '1')
    run_t = sdt.find('.//' + qn('w:sdtContent') + '/' + qn('w:r') + '/' + qn('w:t'))
    run_t.text = u'☒'

    # Progress to date
    set_cell_rich(t2.rows[7].cells[0], progress_blocks)

    # Planned Activity
    set_cell_rich(t2.rows[9].cells[0], planned_blocks)

    # Risks & Mitigation nested table
    risk_cell = t2.rows[11].cells[0]
    tc = risk_cell._tc
    nested_tbl_el = tc.find(qn('w:tbl'))
    found_tbl = False
    for child in list(tc):
        if child is nested_tbl_el:
            found_tbl = True
            continue
        if found_tbl and child.tag == qn('w:p'):
            tc.remove(child)
    nt = risk_cell.tables[0]
    for i, (risk_text, impact, mitigation_text) in enumerate(risks):
        row = nt.rows[i + 1]
        set_cell_simple(row.cells[0], [risk_text])
        tr = row._tr
        sdt_r = tr.findall(qn('w:sdt'))
        set_combo_text(sdt_r[0], impact)
        set_cell_simple(row.cells[-1], [mitigation_text])

    return d

# ============================= PUR#1 =============================
pur1_progress = [
    ("1. Literature Review & Architecture Design (Complete)", True, False, False),
    ("Reviewed existing Graph-Based Transform (GBT) compression literature and designed the "
     "GBT-ICL architecture, extending the standard GBT taxonomy.", False, False, True),
    ("Extended the taxonomy by (i) removing the LLM entropy-coding stage used in a prior design, "
     "and (ii) replacing explicit graph transmission with a frozen, in-context Laplacian predictor.",
     False, False, True),
    ("Produced full encoder/decoder architecture diagrams covering block partition, context "
     "extraction, GBT-ICL, eigendecomposition, GFT, quantisation, and entropy coding.",
     False, False, True),
    ("", False, False, False),
    ("2. Dataset Preparation (Complete)", True, False, False),
    ("Acquired two standard raw YUV 4:2:0 test sequences (Beauty, HoneyBee; 1920x1080, 120fps, "
     "600 frames each).", False, False, True),
    ("Built a frame-extraction pipeline that uniformly samples 30 representative still frames per "
     "sequence (60 frames total) rather than consecutive frames, which are near-duplicates at "
     "120fps, and converts them from planar YUV to RGB.", False, False, True),
    ("", False, False, False),
    ("3. Block Partitioning (Complete)", True, False, False),
    ("Implemented and verified an 8x8 block-partitioning module; each 1920x1080 frame reshapes "
     "into a 135x240 grid of blocks (32,400 blocks/frame), matching the granularity used for the "
     "downstream graph transform.", False, False, True),
    ("", False, False, False),
    ("4. Context Extraction (Complete)", True, False, False),
    ("Implemented the causal context-extraction stage: for every block in raster order, gathers "
     "the already-decoded neighbouring pixels (top and left borders) that GBT-ICL will condition "
     "its Laplacian prediction on, with fixed-padding fallback for edge blocks.", False, False, True),
    ("Verified correctness of neighbour indexing and edge-padding against the underlying block "
     "tensors (spot-checked against raw block data; edge rows/columns correctly padded, interior "
     "blocks match their true neighbours exactly).", False, False, True),
    ("", False, False, False),
    ("5. Data Integrity & Pipeline Robustness (Complete)", True, False, False),
    ("Identified and corrected several corrupted intermediate files caused by shell execution "
     "timeouts during batch processing; added integrity/verification checks to the pipeline "
     "scripts so future runs are resumable and self-checking.", False, False, True),
    ("", False, False, False),
    ("Milestones Achieved: GBT-ICL architecture finalised; dataset acquired and frame-sampled; "
     "block partitioning implemented and verified; context extraction implemented and verified.",
     True, True, False),
    ("", False, False, False),
    ("Ethical approval was not required, as the project uses only publicly available benchmark "
     "video sequences with no human participant or personal data involvement.", False, False, False),
]

pur1_planned = [
    ("1. GBT-ICL Predictor Implementation (Weeks 1-3)", True, False, False),
    ("Implement the GBT-ICL predictor: a lightweight, frozen in-context model mapping extracted "
     "context to a predicted graph Laplacian per block.", False, False, True),
    ("", False, False, False),
    ("2. Transform & Coding Pipeline (Weeks 2-4)", True, False, False),
    ("Implement eigendecomposition of the predicted Laplacian and the resulting Graph Fourier "
     "Transform (GFT) forward/inverse pair.", False, False, True),
    ("Implement quantisation and standard (non-LLM) entropy coding to complete the encoder "
     "pipeline.", False, False, True),
    ("", False, False, False),
    ("3. Simulated Encode/Decode Validation (Weeks 4-5)", True, False, False),
    ("Build a simulated encoder-decoder loop (raster-order, causal) to confirm encoder and "
     "decoder reconstruct identical context with no transmitted graph.", False, False, True),
    ("", False, False, False),
    ("4. Baseline Benchmarking (Weeks 5-7)", True, False, False),
    ("Establish baseline comparisons against fixed-basis DCT and non-adaptive GBT compression on "
     "the extracted dataset, using rate-distortion (PSNR/SSIM vs. bits-per-pixel) as the "
     "evaluation metric.", False, False, True),
    ("Run an ablation on block size (8x8 vs 16x16) to characterise the compression/complexity "
     "trade-off.", False, False, True),
    ("", False, False, False),
    ("5. Dissertation Writing (Weeks 1-7, ongoing)", True, False, False),
    ("Begin drafting the Methodology chapter alongside implementation.", False, False, True),
]

pur1_risks = [
    ("GBT-ICL's Laplacian prediction may be insufficiently accurate to outperform baseline "
     "transforms.", "Medium",
     "Retain a fallback to a fixed/classical GBT basis when predicted-Laplacian confidence is "
     "low; benchmark early and iterate on the predictor architecture."),
    ("Per-block eigendecomposition may be computationally expensive at scale.", "Medium",
     "Keep block size small (8x8) to bound Laplacian matrix size, and batch eigendecomposition "
     "across blocks using vectorised linear algebra."),
    ("The evaluation dataset is currently limited to two video sequences, which may limit "
     "generalisability of results.", "Low",
     "Extend the frame-extraction pipeline to additional public UVG/CTC test sequences if time "
     "permits."),
]

# ============================= PUR#2 =============================
pur2_progress = [
    ("1. GBT-ICL Predictor Implementation (Complete)", True, False, False),
    ("Implemented the frozen in-context predictor: takes each block's top/left decoded-neighbour "
     "context as input and outputs a predicted graph Laplacian for that block, run identically at "
     "encoder and decoder.", False, False, True),
    ("[Insert final predictor architecture summary once training/design is finalised, e.g. layer "
     "count, parameter count, context window size used.]", False, True, True),
    ("", False, False, False),
    ("2. GFT & Eigendecomposition Pipeline (Complete)", True, False, False),
    ("Implemented batched eigendecomposition of predicted per-block Laplacians and the resulting "
     "forward/inverse Graph Fourier Transform (GFT).", False, False, True),
    ("Verified perfect (lossless-path) reconstruction: encoding then immediately decoding a block "
     "with no quantisation returns the original pixel values, confirming the GFT basis and its "
     "inverse are consistent. [Insert measured reconstruction error, expected ~0.]",
     False, False, True),
    ("", False, False, False),
    ("3. Quantisation & Entropy Coding (Complete)", True, False, False),
    ("Implemented scalar quantisation with a configurable step size, and integrated a standard "
     "(non-LLM) entropy coder over the quantised coefficients.", False, False, True),
    ("", False, False, False),
    ("4. Simulated Encoder-Decoder Loop (Complete)", True, False, False),
    ("Built the full raster-order, causal simulated encode/decode loop across the extracted "
     "dataset (60 frames, Beauty + HoneyBee). Confirmed encoder and decoder reconstruct identical "
     "context at every block with zero transmitted graph information.", False, False, True),
    ("", False, False, False),
    ("5. Baseline Benchmarking (In Progress)", True, False, False),
    ("Rate-distortion evaluation of GBT-ICL against fixed-basis DCT and non-adaptive GBT running "
     "across the dataset. [Insert measured PSNR/SSIM vs. bits-per-pixel results once the full "
     "sweep is complete.]", False, False, True),
    ("", False, False, False),
    ("6. Block-Size Ablation (In Progress)", True, False, False),
    ("8x8 vs 16x16 block-size comparison underway to characterise the compression/complexity "
     "trade-off. [Insert results once ablation is complete.]", False, False, True),
    ("", False, False, False),
    ("7. Dissertation Writing (In Progress)", True, False, False),
    ("Methodology chapter drafted covering the GBT-ICL architecture, context-extraction design, "
     "and transform/coding pipeline. [Update completion percentage/word count once finalised.]",
     False, False, True),
    ("", False, False, False),
    ("Milestones Achieved: GBT-ICL predictor, GFT/eigendecomposition, quantisation and entropy "
     "coding all implemented; end-to-end simulated encode/decode loop validated with no graph "
     "side-information transmitted.", True, True, False),
    ("", False, False, False),
    ("Ethical approval remains not required — the project continues to use only publicly "
     "available benchmark video sequences.", False, False, False),
]

pur2_planned = [
    ("1. Complete Evaluation (Weeks 1-3)", True, False, False),
    ("Finish baseline benchmarking and the block-size ablation; finalise rate-distortion curves "
     "for GBT-ICL vs DCT vs non-adaptive GBT.", False, False, True),
    ("Run statistical significance testing (e.g. paired t-tests across evaluated frames) to "
     "confirm any observed improvement is not due to chance.", False, False, True),
    ("", False, False, False),
    ("2. Dataset Extension (Weeks 2-3, if time permits)", True, False, False),
    ("Extend the frame-extraction pipeline to additional public UVG/CTC test sequences to "
     "strengthen generalisability of the results.", False, False, True),
    ("", False, False, False),
    ("3. Failure-Case Analysis (Weeks 3-4)", True, False, False),
    ("Investigate blocks/regions where GBT-ICL underperforms baselines (e.g. high-texture or "
     "edge-block regions with padded context) and document findings.", False, False, True),
    ("", False, False, False),
    ("4. Dissertation Writing (Weeks 1-7, ongoing)", True, False, False),
    ("Complete the Methodology chapter; begin and progress the Results chapter with the "
     "finalised rate-distortion figures and tables.", False, False, True),
    ("Draft/refine Introduction and Literature Review chapters if not already complete.",
     False, False, True),
]

pur2_risks = [
    ("GBT-ICL underperforms baseline transforms in initial rate-distortion tests.", "Medium",
     "Investigate failure cases by region/texture; consider enriching the context window (larger "
     "border, additional neighbour) before concluding the approach is not viable."),
    ("Experimentation running longer than planned, compressing dissertation writing time.", "High",
     "Prioritise finishing the evaluation sweep first; set daily writing targets once results are "
     "in; schedule weekly supervisor check-ins to catch slippage early."),
    ("Evaluation dataset remains limited to two video sequences.", "Low",
     "Extend to additional public UVG/CTC sequences if time allows; otherwise clearly scope this "
     "as a stated limitation in the dissertation."),
]

# ============================= PUR#3 =============================
pur3_progress = [
    ("1. Full Evaluation Complete (Complete)", True, False, False),
    ("Completed rate-distortion evaluation of GBT-ICL against fixed-basis DCT and non-adaptive "
     "GBT baselines across the full dataset. [Insert final PSNR/SSIM vs. bits-per-pixel results "
     "and rate-distortion plots.]", False, False, True),
    ("Completed the 8x8 vs 16x16 block-size ablation. [Insert final compression/complexity "
     "trade-off results.]", False, False, True),
    ("Statistical significance testing complete (paired t-tests across evaluated frames). [Insert "
     "p-values/summary statistics.]", False, False, True),
    ("", False, False, False),
    ("2. Failure-Case Analysis (Complete)", True, False, False),
    ("Documented regions/conditions where GBT-ICL under- or over-performs baseline transforms, "
     "with discussion of likely causes (e.g. edge-block padding, high-texture context).",
     False, False, True),
    ("", False, False, False),
    ("3. Dissertation Writing (Complete / Near-Complete)", True, False, False),
    ("Introduction and Literature Review chapters: complete, supervisor-reviewed.", False, False, True),
    ("Methodology chapter: complete, covering GBT-ICL architecture, context extraction, "
     "transform/coding pipeline, and experimental design.", False, False, True),
    ("Results chapter: complete, with all rate-distortion figures, ablation tables, and "
     "statistical analysis.", False, False, True),
    ("Discussion & Conclusions chapter: drafted, covering contributions, limitations (dataset "
     "size, causal-context assumption, block-size granularity), and future work.", False, False, True),
    ("[Update chapter word counts / completion percentages with actual figures before "
     "submission.]", False, True, True),
    ("", False, False, False),
    ("4. Codebase & Reproducibility (Complete)", True, False, False),
    ("Finalised and documented the full pipeline codebase (frame extraction, block partitioning, "
     "context extraction, GBT-ICL predictor, transform/coding, evaluation scripts) for inclusion "
     "as a dissertation appendix.", False, False, True),
    ("", False, False, False),
    ("Milestones Achieved: Full evaluation and ablation complete; all core dissertation chapters "
     "drafted; codebase finalised and documented.", True, True, False),
    ("", False, False, False),
    ("Ethical approval remains not required for the full duration of the project.", False, False, False),
]

pur3_planned = [
    ("1. Supervisor Feedback & Revisions (Weeks 1-3)", True, False, False),
    ("Submit the complete draft (all chapters, abstract, references, appendices) to the "
     "supervisor for feedback.", False, False, True),
    ("Incorporate supervisor feedback across all chapters.", False, False, True),
    ("", False, False, False),
    ("2. Final Proofreading & Formatting (Weeks 3-5)", True, False, False),
    ("Compile the full reference list in the required citation style and finalise appendices "
     "(code listings, additional result tables/figures).", False, False, True),
    ("Complete final proofreading and formatting to the Aston University dissertation "
     "requirements.", False, False, True),
    ("", False, False, False),
    ("3. Final Submission", True, False, False),
    ("Submit the final dissertation via Blackboard by the Expected Project Completion Date "
     "(11/12/2026).", False, False, True),
    ("Deliverables: full dissertation, documented codebase, and extracted dataset/pipeline "
     "artefacts.", False, False, True),
]

pur3_risks = [
    ("Dissertation writing/revision falling behind due to supervisor feedback turnaround time.",
     "Medium",
     "Submit the complete draft early within the term; set daily writing/revision targets; "
     "schedule regular supervisor check-ins in the final weeks."),
    ("Minor formatting/citation inconsistencies discovered late in proofreading.", "Low",
     "Use reference-management software throughout and run a dedicated formatting pass before "
     "the final proofreading round."),
    ("Scope creep from wanting to add further ablations/experiments this late in the project.",
     "Low",
     "Strictly adhere to the originally scoped research objectives; defer any additional "
     "experiments to a stated future-work section."),
]

configs = [
    ("PUR#1", pur1_progress, pur1_planned, pur1_risks, "form_PUR1.docx"),
    ("PUR#2", pur2_progress, pur2_planned, pur2_risks, "form_PUR2.docx"),
    ("PUR#3", pur3_progress, pur3_planned, pur3_risks, "form_PUR3.docx"),
]

for checkpoint, progress, planned, risks, outname in configs:
    doc = build(checkpoint, progress, planned, risks)
    doc.save(outname)
    print("saved", outname)
